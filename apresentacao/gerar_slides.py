#!/usr/bin/env python3
"""Gera os slides da apresentação (PPTX editável + PDF) do Grupo 2 – Agenda de Serviços.
Descrição única das telas, renderizada para PowerPoint (python-pptx) e PDF (reportlab).
Uso: python gerar_slides.py  ( gera Agenda-de-Servicos-Grupo2.pptx / .pdf na mesma pasta )
"""
import html
import os

# ---------- constantes ----------
W, H = 13.333, 7.5  # polegadas (16:9)

NAVY = (0, 29, 51)
NAVY2 = (0, 51, 83)
PRIMARY = (0, 99, 154)
PRIMARY_D = (0, 75, 117)
SKY = (206, 229, 255)
SKY2 = (213, 228, 246)
SKY_SOFT = (235, 244, 253)
PURPLE = (104, 88, 122)
PURPLE_L = (239, 220, 255)
BG = (250, 252, 255)
INK = (26, 28, 30)
MUTED = (81, 96, 111)
BORDER = (203, 217, 232)
WHITE = (255, 255, 255)
GREEN = (46, 125, 50); GREEN_BG = (232, 245, 233)
ORANGE = (230, 81, 0); ORANGE_BG = (255, 243, 224)
BLUE_T = (2, 136, 209); BLUE_BG = (225, 245, 254)
RED = (198, 40, 40); RED_BG = (255, 235, 238)
GRAY_BG = (238, 243, 248)

A = {"l": 0, "c": 1, "r": 2}  # alinhamento reportlab

# ---------- backend PowerPoint ----------
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

PA = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}

class PptxCanvas:
    kind = "pptx"

    def __init__(self, path):
        self.path = path
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.slide = None

    def begin(self, bg=BG):
        self.slide = self.prs.slides.add_slide(self.blank)
        self.rect(0, 0, W, H, bg)

    def end(self):
        pass

    def note(self, text):
        ns = self.slide.notes_slide
        ns.notes_text_frame.text = text

    @staticmethod
    def _style(shape, fill=None, line=None, lw=1.0):
        shape.shadow.inherit = False
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(*fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = RGBColor(*line); shape.line.width = Pt(lw)

    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.0):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius > 0 else MSO_SHAPE.RECTANGLE
        s = self.slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius > 0:
            try:
                s.adjustments[0] = max(0.0, min(0.5, radius / max(min(w, h), 0.01) * 0.5))
            except Exception:
                pass
        self._style(s, fill, line, lw)
        return s

    def oval(self, x, y, w, h, fill=None, line=None, lw=1.0):
        s = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
        self._style(s, fill, line, lw)
        return s

    def line(self, x1, y1, x2, y2, color=BORDER, lw=1.0):
        s = self.slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        s.line.color.rgb = RGBColor(*color); s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def arrow(self, x1, y, x2, color=PRIMARY, lw=2.0):
        self.line(x1, y, x2 - 0.13, y, color, lw)
        s = self.slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                        Inches(x2 - 0.18), Inches(y - 0.09), Inches(0.18), Inches(0.18))
        s.rotation = 90
        self._style(s, color, None)

    def text(self, x, y, w, h, txt, size=14, bold=False, color=INK, align="l",
             leading=1.08, mono=False, anchor="t", space_after=2):
        box = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if anchor == "m" else MSO_ANCHOR.TOP
        lines = str(txt).split("\n")
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PA[align]
            p.line_spacing = leading
            p.space_after = Pt(space_after)
            r = p.add_run(); r.text = ln
            r.font.size = Pt(size); r.font.bold = bold
            r.font.name = "Consolas" if mono else "Calibri"
            r.font.color.rgb = RGBColor(*color)
        return box

    def save(self):
        cp = self.prs.core_properties
        cp.title = "Agenda de Serviços — Grupo 2"
        cp.author = "Grupo 2"
        cp.subject = "Trabalho Prático – Desenvolvimento Web / PWA"
        self.prs.save(self.path)

# ---------- backend PDF ----------
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import Paragraph

FONT_DIR = "/usr/share/fonts/truetype/dejavu"
try:
    pdfmetrics.registerFont(TTFont("Sans", f"{FONT_DIR}/DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("SansB", f"{FONT_DIR}/DejaVuSans-Bold.ttf"))
    pdfmetrics.registerFont(TTFont("Mono", f"{FONT_DIR}/DejaVuSansMono.ttf"))
    pdfmetrics.registerFont(TTFont("MonoB", f"{FONT_DIR}/DejaVuSansMono-Bold.ttf"))
except Exception:
    pass

def hc(c):
    return HexColor("#%02x%02x%02x" % c)

class PdfCanvas:
    kind = "pdf"

    def __init__(self, path):
        self.path = path
        self.c = pdf_canvas.Canvas(path, pagesize=(W * 72, H * 72))
        self.c.setTitle("Agenda de Serviços — Grupo 2")
        self.c.setAuthor("Grupo 2")

    def begin(self, bg=BG):
        self.rect(0, 0, W, H, bg)

    def end(self):
        self.c.showPage()

    def note(self, text):
        pass  # notas do apresentador ficam no PPTX

    def _paint(self, fill, line, lw):
        if fill is not None:
            self.c.setFillColor(hc(fill))
        if line is not None:
            self.c.setStrokeColor(hc(line)); self.c.setLineWidth(lw)
        return (1 if fill is not None else 0), (1 if line is not None else 0)

    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.0):
        X, Y = x * 72, (H - y - h) * 72
        f, s = self._paint(fill, line, lw)
        if radius > 0:
            self.c.roundRect(X, Y, w * 72, h * 72, radius * 72, stroke=s, fill=f)
        else:
            self.c.rect(X, Y, w * 72, h * 72, stroke=s, fill=f)

    def oval(self, x, y, w, h, fill=None, line=None, lw=1.0):
        f, s = self._paint(fill, line, lw)
        self.c.ellipse(x * 72, (H - y - h) * 72, (x + w) * 72, (H - y) * 72, stroke=s, fill=f)

    def line(self, x1, y1, x2, y2, color=BORDER, lw=1.0):
        self.c.setStrokeColor(hc(color)); self.c.setLineWidth(lw)
        self.c.line(x1 * 72, (H - y1) * 72, x2 * 72, (H - y2) * 72)

    def arrow(self, x1, y, x2, color=PRIMARY, lw=2.0):
        self.line(x1, y, x2 - 0.13, y, color, lw)
        self.c.setFillColor(hc(color)); self.c.setStrokeColor(hc(color))
        p = self.c.beginPath()
        p.moveTo(x2 * 72, (H - y) * 72)
        p.lineTo((x2 - 0.18) * 72, (H - y - 0.09) * 72)
        p.lineTo((x2 - 0.18) * 72, (H - y + 0.09) * 72)
        p.close()
        self.c.drawPath(p, stroke=0, fill=1)

    def text(self, x, y, w, h, txt, size=14, bold=False, color=INK, align="l",
             leading=1.08, mono=False, anchor="t", space_after=2):
        fname = ("MonoB" if bold else "Mono") if mono else ("SansB" if bold else "Sans")
        try:
            pdfmetrics.getFont(fname)
        except Exception:
            fname = "Helvetica-Bold" if bold else "Helvetica"
        al = {"l": TA_LEFT, "c": TA_CENTER, "r": TA_RIGHT}[align]
        st = ParagraphStyle("s", fontName=fname, fontSize=size, leading=size * leading + space_after,
                            textColor=hc(color), alignment=al, spaceAfter=0)
        esc = html.escape(str(txt)).replace("\n", "<br/>")
        p = Paragraph(esc, st)
        pw, ph = p.wrap(w * 72, h * 72)
        top = (H - y) * 72
        yy = top - ph if anchor == "t" else top - (h * 72 + ph) / 2
        p.drawOn(self.c, x * 72, yy)

    def save(self):
        self.c.save()

# ---------- helpers de desenho ----------
def pill(c, x, y, w, h, txt, fill, color, size=11, bold=True, line=None):
    c.rect(x, y, w, h, fill, line=line, radius=h / 2)
    c.text(x + 0.06, y, w - 0.12, h, txt, size=size, bold=bold, color=color, align="c", anchor="m")

def chip_status(c, x, y, w, txt, bgc, tc):
    pill(c, x, y, w, 0.34, txt, bgc, tc, size=10.5)

def card(c, x, y, w, h, fill=WHITE, line=BORDER, radius=0.14):
    c.rect(x, y, w, h, fill, line=line, lw=1.0, radius=radius)

def num_badge(c, x, y, d, n, fill=PRIMARY, color=WHITE, size=15):
    c.oval(x, y, d, d, fill)
    c.text(x, y + 0.01, d, d - 0.02, str(n), size=size, bold=True, color=color, align="c", anchor="m")

def header(c, n, section, title, subtitle=None):
    pill(c, 0.55, 0.38, max(1.55, 0.34 + 0.104 * len(section)), 0.36, section, SKY, NAVY, size=10)
    c.text(11.6, 0.38, 1.18, 0.36, f"{n:02d} / 14", size=11, bold=True, color=MUTED, align="r", anchor="m")
    c.text(0.55, 0.84, 12.2, 0.62, title, size=28, bold=True, color=NAVY)
    yy = 1.46
    if subtitle:
        c.text(0.55, 1.48, 12.2, 0.4, subtitle, size=14, color=MUTED)
        yy = 1.88
    c.rect(0.55, yy, 0.85, 0.055, PRIMARY, radius=0.03)
    c.text(0.55, 7.12, 6.0, 0.28, "Grupo 2  •  Agenda de Serviços", size=9.5, color=MUTED)
    c.text(9.3, 7.12, 3.48, 0.28, "Trabalho Prático – Desenvolvimento Web / PWA", size=9.5, color=MUTED, align="r")
    return yy + 0.22

def phone(c, x, y, w, h, title, variant):
    c.rect(x, y, w, h, NAVY, radius=0.30)
    c.rect(x + 0.09, y + 0.09, w - 0.18, h - 0.18, WHITE, radius=0.24)
    c.rect(x + 0.09, y + 0.09, w - 0.18, 0.72, PRIMARY, radius=0.24)
    c.rect(x + 0.09, y + 0.55, w - 0.18, 0.26, PRIMARY)  # fecha cantos inferiores do topo
    pill(c, x + w/2 - 0.30, y + 0.20, 0.60, 0.13, "", NAVY, NAVY, size=6)
    c.text(x + 0.26, y + 0.30, w - 0.52, 0.4, title, size=13, bold=True, color=WHITE)
    cx, cw, cy = x + 0.26, w - 0.52, y + 1.02
    compact = h < 5.30
    if variant == "agenda":
        pill(c, cx, cy, cw, 0.30 if compact else 0.34, "Hoje  •  Terça-feira", SKY, NAVY, size=9.5)
        items = [("09:00", "Ana • Corte feminino", "Confirmado", BLUE_BG, BLUE_T),
                 ("10:30", "João • barba", "Pendente", ORANGE_BG, ORANGE),
                 ("11:15", "Maria • Coloração", "Concluído", GREEN_BG, GREEN)]
        yy = cy + (0.44 if compact else 0.52)
        for hora, nome, st, b, t in items:
            card(c, cx, yy, cw, 0.76 if compact else 0.86, fill=GRAY_BG, line=None, radius=0.10)
            c.text(cx + 0.12, yy + 0.10, 0.85, 0.3, hora, size=10.5, bold=True, color=PRIMARY)
            c.text(cx + 0.12, yy + 0.40, cw - 0.2, 0.3, nome, size=9.5, color=INK)
            chip = {"Confirmado": 1.05, "Pendente": 0.95, "Concluído": 0.95}[st]
            pill(c, cx + cw - chip - 0.10, yy + 0.08, chip, 0.30, st, b, t, size=8)
            yy += 0.86 if compact else 1.00
        plus = 0.40 if compact else 0.52
        c.oval(cx + cw - plus - 0.03, yy + 0.02, plus, plus, PRIMARY)
        c.text(cx + cw - plus - 0.03, yy + 0.03, plus, plus - 0.04, "+", size=16 if compact else 19, bold=True, color=WHITE, align="c", anchor="m")
    elif variant == "clientes":
        card(c, cx, cy, cw, 0.42, fill=GRAY_BG, line=None, radius=0.10)
        c.text(cx + 0.14, cy + 0.05, cw - 0.2, 0.32, "Buscar cliente...", size=9.5, color=MUTED, anchor="m")
        pessoas = [("A", "Ana", "(62) 99999-1020"), ("J", "João", "(62) 98888-7788"),
                   ("M", "Maria", "(62) 97777-4455"), ("P", "Pedro Lima", "(62) 96666-7788")]
        yy = cy + (0.58 if compact else 0.62)
        avatar = 0.48 if compact else 0.52
        for ini, nome, tel in pessoas:
            c.oval(cx, yy, avatar, avatar, SKY)
            c.text(cx, yy + 0.02, avatar, avatar - 0.04, ini, size=12 if compact else 13, bold=True, color=NAVY, align="c", anchor="m")
            c.text(cx + avatar + 0.14, yy - 0.02, cw - avatar - 0.18, 0.3, nome, size=10.5, bold=True, color=INK)
            c.text(cx + avatar + 0.14, yy + 0.27, cw - avatar - 0.18, 0.3, tel, size=9, color=MUTED)
            yy += 0.70 if compact else 0.78
    elif variant == "servicos":
        itens = [("Corte feminino", "45 min  •  R$ 70,00"), ("barba", "30 min  •  R$ 35,00"),
                 ("Coloração", "90 min  •  R$ 140,00"), ("Troca de óleo", "60 min  •  R$ 180,00")]
        yy = cy
        for nome, det in itens:
            card(c, cx, yy, cw, 0.70 if compact else 0.78, fill=GRAY_BG, line=None, radius=0.10)
            c.text(cx + 0.14, yy + 0.08, cw - 0.2, 0.3, nome, size=11, bold=True, color=INK)
            c.text(cx + 0.14, yy + (0.36 if compact else 0.41), cw - 0.2, 0.3, det, size=9.5, color=PRIMARY)
            yy += 0.82 if compact else 0.94
    elif variant == "form":
        campos = [("Cliente", "Ana"), ("Serviço", "Corte feminino (45 min)"),
                  ("Data", "20/10/2026"), ("Horário", "09:00 – 09:45")]
        yy = cy
        for rot, val in campos:
            c.text(cx, yy, cw, 0.26, rot, size=9, bold=True, color=MUTED)
            card(c, cx, yy + 0.23, cw, 0.40, fill=GRAY_BG, line=BORDER, radius=0.09)
            c.text(cx + 0.12, yy + 0.25, cw - 0.2, 0.36, val, size=10, color=INK, anchor="m")
            yy += 0.74
        pill(c, cx, yy + 0.02, cw, 0.44, "Salvar agendamento", PRIMARY, WHITE, size=10.5)
        c.text(cx, yy + 0.56, cw, 0.3, "✓ Horário livre — sem conflito", size=9, bold=True, color=GREEN)
    # barra inferior
    by = y + h - 0.52
    c.line(x + 0.28, by, x + w - 0.28, by, BORDER, 0.75)
    for i, rot in enumerate(["Agenda", "Clientes", "Serviços"]):
        xx = x + 0.30 + i * (w - 0.60) / 3
        active = (variant == rot.lower()) or (variant == "form" and i == 0)
        c.text(xx, by + 0.08, (w - 0.60) / 3, 0.26, rot, size=8, bold=active,
               color=PRIMARY if active else MUTED, align="c")

# ---------- slides ----------
def s01(c):
    c.begin(NAVY)
    c.note("TEMPO: 0:00–0:45 | APRESENTADOR(A) 1\n"
           "• Bom dia/boa tarde. Somos o Grupo 2 e nosso tema é Agenda de Serviços para barbearia/oficina.\n"
           "• Em uma frase: criamos um app para cadastrar clientes e serviços, agendar sem conflito de horário e lembrar do atendimento.\n"
           "• Avise a estrutura: problema, solução, tecnologias, 4 etapas, 3 desafios, o que funcionou, divergências em relação ao PWA, próximos passos e lições.\n"
           "ANTES DE APRESENTAR: substituir os nomes/turma/professor marcados entre colchetes neste slide.")
    c.oval(9.4, -2.2, 6.2, 6.2, NAVY2)
    c.oval(-2.4, 4.6, 5.4, 5.4, PRIMARY_D)
    c.oval(6.5, 5.5, 2.6, 2.6, PURPLE)
    pill(c, 0.65, 0.55, 4.45, 0.42, "TRABALHO PRÁTICO • DESENVOLVIMENTO WEB / PWA", None, SKY, size=10, line=PRIMARY)
    c.text(0.65, 1.25, 6.6, 2.2, "Agenda de\nServiços", size=52, bold=True, color=WHITE, leading=1.02)
    c.text(0.65, 3.25, 6.3, 1.0, "Agendamento inteligente para barbearias,\nsalões, oficinas e prestadores de serviços",
           size=18, color=SKY, leading=1.15)
    c.rect(0.65, 4.35, 1.1, 0.06, SKY, radius=0.03)
    c.text(0.65, 4.62, 6.6, 0.45, "Grupo 2 — 4 integrantes", size=17, bold=True, color=WHITE)
    c.text(0.65, 5.12, 6.6, 0.85,
           "Integrantes: [Nome 1]  •  [Nome 2]  •  [Nome 3]  •  [Nome 4]\nTurma/Disciplina: [preencher]  •  Professor(a): [preencher]",
           size=12.5, color=SKY2, leading=1.25)
    pill(c, 0.65, 6.25, 3.05, 0.44, "Apresentação • 15 minutos", PRIMARY, WHITE, size=11)
    phone(c, 7.85, 0.85, 2.45, 5.85, "Agenda", "agenda")
    phone(c, 10.35, 1.25, 2.35, 5.45, "Novo agendamento", "form")

def s02(c):
    c.begin(BG)
    c.note("TEMPO: 0:45–1:45 | APRESENTADOR(A) 1\n"
           "• Situe a atividade: a especificação pede um PWA hospedado; o Grupo 2 ficou com Agenda de Serviços.\n"
           "• Sejam transparentes logo no início: o que foi construído é um app Android nativo funcional. Não vamos chamar de PWA o que não é.\n"
           "• Promessa da apresentação: demonstrar a solução e, no fim, mostrar com clareza onde ela diverge da especificação e o plano para adequar.\n"
           "TRANSIÇÃO: 'Primeiro, qual problema o app resolve?'")
    y = header(c, 2, "CONTEXTO", "O ponto de partida", "O que foi pedido — e o que trazemos para apresentar")
    card(c, 0.55, y, 5.95, 3.55)
    num_badge(c, 0.85, y + 0.28, 0.5, "!")
    c.text(1.5, y + 0.30, 4.7, 0.5, "A especificação do trabalho", size=18, bold=True, color=NAVY, anchor="m")
    c.text(0.85, y + 1.02, 5.35, 2.35,
           "• Desenvolver um aplicativo web em formato PWA, hospedado em nuvem\n"
           "• Grupo 2: App Agenda de Serviços (barbearia/oficina) — cadastro de cliente, horários e status, com calendário\n"
           "• Entregáveis: link público do app + link do repositório GitHub",
           size=14.5, color=INK, leading=1.22, space_after=6)
    card(c, 6.85, y, 5.93, 3.55, fill=SKY_SOFT, line=SKY)
    num_badge(c, 7.15, y + 0.28, 0.5, "✓")
    c.text(7.8, y + 0.30, 4.7, 0.5, "O que construímos", size=18, bold=True, color=NAVY, anchor="m")
    c.text(7.15, y + 1.02, 5.35, 2.35,
           "• Um app funcional de agenda de serviços, com o tema e as funções do Grupo 2\n"
           "• Construído como Android nativo (Kotlin), não como página web\n"
           "• Nesta apresentação, mostramos o que atende, onde diverge do PWA e como corrigir o rumo",
           size=14.5, color=INK, leading=1.22, space_after=6)
    c.rect(0.55, y + 3.85, 12.23, 0.78, NAVY, radius=0.14)
    c.text(0.9, y + 3.85, 11.6, 0.78, "Nosso compromisso nesta apresentação: demonstrar a solução, as decisões técnicas e o aprendizado — sem esconder as pendências.",
           size=14.5, bold=True, color=WHITE, align="c", anchor="m")

def s03(c):
    c.begin(BG)
    c.note("TEMPO: 1:45–2:45 | APRESENTADOR(A) 1\n"
           "• Descreva a rotina real: agenda no papel/WhatsApp, telefone do cliente perdido, dois clientes no mesmo horário, falta sem aviso.\n"
           "• Reforce o público: pequeno negócio, muitas vezes sem sistema pago e com internet instável — por isso 'funcionar offline' importa.\n"
           "TRANSIÇÃO: 'Foi para atacar essas três dores que desenhamos o app.'")
    y = header(c, 3, "O PROBLEMA", "Agenda manual gera retrabalho", "Três dores de quem atende com hora marcada")
    dores = [
        ("Informação espalhada", "Clientes, preços e horários ficam no caderno, na cabeça e nas conversas do WhatsApp.",
         "Resultado: difícil consultar, buscar e manter o histórico."),
        ("Conflito de horário", "Sem validar a duração de cada serviço, dois atendimentos acabam sobrepostos.",
         "Resultado: espera, atraso em cadeia e cliente insatisfeito."),
        ("Esquecimento", "Nem o profissional nem o cliente recebem um lembrete perto do atendimento.",
         "Resultado: faltas (no-show) e horários vagos que poderiam render."),
    ]
    for i, (t, b, r) in enumerate(dores):
        x = 0.55 + i * 4.18
        card(c, x, y, 3.92, 3.30)
        c.rect(x, y, 3.92, 0.10, [ORANGE, RED, PURPLE][i], radius=0.05)
        num_badge(c, x + 0.28, y + 0.32, 0.52, i + 1, fill=[ORANGE, RED, PURPLE][i])
        c.text(x + 0.95, y + 0.34, 2.75, 0.5, t, size=16.5, bold=True, color=NAVY, anchor="m")
        c.text(x + 0.28, y + 1.05, 3.36, 1.35, b, size=13.5, color=INK, leading=1.2)
        c.text(x + 0.28, y + 2.45, 3.36, 0.7, r, size=12.5, bold=True, color=MUTED, leading=1.18)
    c.rect(0.55, y + 3.60, 12.23, 0.95, SKY_SOFT, line=SKY, radius=0.14)
    c.text(0.85, y + 3.60, 11.65, 0.95,
           "Público-alvo: barbearias, salões de beleza, oficinas e prestadores autônomos que precisam controlar atendimentos pelo celular — mesmo sem internet.",
           size=14.5, bold=True, color=NAVY, align="c", anchor="m")

def s04(c):
    c.begin(BG)
    c.note("TEMPO: 2:45–4:00 | APRESENTADOR(A) 2\n"
           "• Passe pelas 6 capacidades ligando cada uma a uma dor do slide anterior.\n"
           "• Destaque o diferencial: o sistema calcula início/fim pela duração do serviço e impede sobreposição antes de salvar.\n"
           "• Mostre o fluxo de status: todo atendimento nasce Pendente e evolui até Concluído ou Cancelado.\n"
           "TRANSIÇÃO: 'Vamos ver isso nas telas.'")
    y = header(c, 4, "A SOLUÇÃO", "O que o app entregue faz", "Funcionalidades ligadas diretamente às três dores")
    feats = [
        ("Clientes", "Cadastro com telefone/WhatsApp, busca e edição.", PRIMARY),
        ("Serviços", "Preço em R$ e duração em minutos de cada atendimento.", PRIMARY),
        ("Agendamentos", "Cliente + serviço + data e horário, com observações.", PRIMARY),
        ("Sem conflito de horário", "Validação de sobreposição que bloqueia choque de agenda.", PURPLE),
        ("Visão da agenda", "Lista por dia, busca em tempo real e filtros por status.", PURPLE),
        ("Lembretes locais", "Notificação 15 minutos antes, mesmo com o app fechado.", PURPLE),
    ]
    for i, (t, b, acc) in enumerate(feats):
        x = 0.55 + (i % 3) * 4.18
        yy = y + (i // 3) * 1.62
        card(c, x, yy, 3.92, 1.44)
        c.oval(x + 0.24, yy + 0.24, 0.46, 0.46, acc)
        c.text(x + 0.24, yy + 0.25, 0.46, 0.44, "✓", size=15, bold=True, color=WHITE, align="c", anchor="m")
        c.text(x + 0.86, yy + 0.20, 2.9, 0.42, t, size=15.5, bold=True, color=NAVY)
        c.text(x + 0.86, yy + 0.62, 2.88, 0.72, b, size=12, color=MUTED, leading=1.15)
    yy = y + 3.45
    c.text(0.55, yy, 3.4, 0.4, "Ciclo de vida do atendimento:", size=13.5, bold=True, color=NAVY, anchor="m")
    xs = 4.05
    for j, (rot, b, t) in enumerate([("Pendente", ORANGE_BG, ORANGE), ("Confirmado", BLUE_BG, BLUE_T), ("Concluído", GREEN_BG, GREEN)]):
        pill(c, xs, yy, 1.62, 0.42, rot, b, t, size=12)
        ligacao = "ou" if j == 2 else "→"
        c.text(xs + 1.62, yy, 0.42, 0.42, ligacao, size=12 if j == 2 else 15, bold=True, color=MUTED, align="c", anchor="m")
        xs += 2.04
    pill(c, xs, yy, 1.62, 0.42, "Cancelado", RED_BG, RED, size=12)
    c.text(0.55, yy + 0.62, 12.2, 0.4, "Tudo salvo localmente no aparelho: o app funciona 100% offline (detalhes técnicos nos próximos slides).",
           size=13, color=MUTED, align="c")

def s05(c):
    c.begin(BG)
    c.note("TEMPO: 4:00–5:15 | APRESENTADOR(A) 2\n"
           "ROTEIRO DE DEMONSTRAÇÃO AO VIVO (ensaiar com dados já cadastrados):\n"
           "1) Abrir a Agenda do dia; 2) cadastrar um cliente; 3) cadastrar/confirmar um serviço;\n"
           "4) criar agendamento válido; 5) tentar criar outro sobreposto e mostrar o bloqueio; 6) mudar um status.\n"
           "Os celulares deste slide são ilustrações fiéis às telas — se tiverem prints reais, substituam as imagens no PPTX.")
    y = header(c, 5, "DEMONSTRAÇÃO", "Do cadastro ao lembrete", "As telas principais e o fluxo que vamos mostrar ao vivo")
    phone(c, 0.55, y - 0.05, 2.30, 5.00, "Agenda", "agenda")
    phone(c, 3.05, y - 0.05, 2.30, 5.00, "Clientes", "clientes")
    phone(c, 5.55, y - 0.05, 2.30, 5.00, "Serviços", "servicos")
    card(c, 8.15, y - 0.05, 4.63, 5.00, fill=NAVY, line=None)
    c.text(8.45, y + 0.22, 4.05, 0.5, "Fluxo de uso", size=18, bold=True, color=WHITE)
    passos = [
        ("1", "Cadastrar cliente e serviço", "Base reutilizável em todos os agendamentos."),
        ("2", "Escolher data e horário", "O fim é calculado pela duração do serviço."),
        ("3", "Validar conflito", "Sobreposição é bloqueada antes de salvar."),
        ("4", "Acompanhar a agenda", "Buscar, filtrar por status e atualizar."),
        ("5", "Receber o lembrete", "Notificação 15 min antes do atendimento."),
    ]
    yy = y + 0.85
    for n, t, b in passos:
        num_badge(c, 8.45, yy, 0.44, n, fill=PRIMARY, size=13)
        c.text(9.02, yy - 0.04, 3.55, 0.34, t, size=13.5, bold=True, color=WHITE)
        c.text(9.02, yy + 0.30, 3.55, 0.42, b, size=11.5, color=SKY2, leading=1.12)
        yy += 0.82

def s06(c):
    c.begin(BG)
    c.note("TEMPO: 5:15–6:30 | APRESENTADOR(A) 3\n"
           "• Apresente a stack por camadas (tela, arquitetura, dados, sistema) e o motivo de cada escolha.\n"
           "• Compose: interface declarativa, menos código de tela. Room: SQLite com consultas tipadas. Flow/StateFlow: a tela se atualiza sozinha quando o banco muda.\n"
           "• Feche sendo direto: é uma stack Android nativa — é justamente aqui que nasce a divergência com o PWA, que detalharemos no slide 11.")
    y = header(c, 6, "TECNOLOGIAS", "A base técnica do app", "Android nativo moderno: Kotlin + Jetpack Compose + Room")
    grupos = [
        ("Interface", PRIMARY, "Kotlin 2.1.10\nJetpack Compose + Material Design 3\nNavigation Compose (barra inferior)"),
        ("Arquitetura", PURPLE, "MVVM + Repository Pattern\nCoroutines, Flow e StateFlow\n4 ViewModels (1 por tela)"),
        ("Dados offline", PRIMARY, "Room 2.8.5 sobre SQLite\nKSP (geração de código do Room)\nConsultas reativas com Flow"),
        ("Recursos de sistema", PURPLE, "AlarmManager + NotificationManager\nBroadcastReceivers (alarme e reinício)\nApp Shortcuts no ícone do app"),
    ]
    for i, (t, acc, b) in enumerate(grupos):
        x = 0.55 + (i % 2) * 4.05
        yy = y + (i // 2) * 1.98
        card(c, x, yy, 3.80, 1.82)
        c.rect(x, yy, 0.10, 1.82, acc, radius=0.05)
        c.text(x + 0.30, yy + 0.16, 3.3, 0.4, t, size=16, bold=True, color=NAVY)
        c.text(x + 0.30, yy + 0.60, 3.32, 1.1, b, size=12.5, color=INK, leading=1.22, space_after=3)
    card(c, 8.85, y, 3.93, 3.80, fill=SKY_SOFT, line=SKY)
    c.text(9.12, y + 0.20, 3.4, 0.45, "Por que essas escolhas?", size=16.5, bold=True, color=NAVY)
    c.text(9.12, y + 0.72, 3.40, 2.25,
           "• UI declarativa acelera a construção das telas\n"
           "• Banco local = offline de verdade, sem servidor\n"
           "• Estado reativo mantém agenda sempre atual\n"
           "• Acesso direto a alarmes e notificações do sistema",
           size=12.5, color=INK, leading=1.22, space_after=5)
    pill(c, 9.12, y + 3.05, 3.40, 0.52, "Android nativo ≠ PWA web", ORANGE_BG, ORANGE, size=11)

def s07(c):
    c.begin(BG)
    c.note("TEMPO: 6:30–7:45 | APRESENTADOR(A) 3\n"
           "• Siga as setas: a tela nunca fala direto com o banco — ViewModel pede ao Repository, que usa o DAO do Room.\n"
           "• Explique as 3 tabelas e que o agendamento guarda início e FIM (início + duração do serviço).\n"
           "• Mostre a regra de conflito no quadro escuro: dois intervalos se sobrepõem quando o novo início é antes do fim do existente E o novo fim é depois do início do existente. Cancelados e o próprio registro em edição ficam fora.")
    y = header(c, 7, "TECNOLOGIAS", "Arquitetura MVVM e dados offline", "Separação de responsabilidades do clique até o banco SQLite")
    boxes = [("Telas Compose", "Agenda • Clientes\nServiços • Formulário"),
             ("ViewModels", "Estado com StateFlow\nValidações da tela"),
             ("Repositories", "Regras de acesso\naos dados"),
             ("DAOs Room", "Consultas SQL\nCliente • Serviço • Agenda"),
             ("SQLite (Room)", "Banco local no aparelho\n100% offline")]
    bw = 2.28
    for i, (t, b) in enumerate(boxes):
        x = 0.55 + i * (bw + 0.21)
        card(c, x, y, bw, 1.52, fill=NAVY if i == 4 else WHITE, line=None if i == 4 else BORDER)
        c.text(x + 0.16, y + 0.16, bw - 0.32, 0.42, t, size=14.5, bold=True, color=WHITE if i == 4 else NAVY, align="c")
        c.text(x + 0.16, y + 0.62, bw - 0.32, 0.8, b, size=11.5, color=SKY2 if i == 4 else MUTED, align="c", leading=1.18)
        if i < 4:
            c.arrow(x + bw + 0.02, y + 0.76, x + bw + 0.19)
    yy = y + 1.85
    card(c, 0.55, yy, 6.45, 2.75)
    c.text(0.83, yy + 0.20, 5.9, 0.42, "Modelo de dados (3 entidades)", size=16, bold=True, color=NAVY)
    c.text(0.83, yy + 0.68, 5.95, 1.95,
           "Cliente — nome e telefone/WhatsApp\n"
           "Serviço — nome, preço (R$) e duração (min)\n"
           "Agendamento — cliente + serviço (chaves estrangeiras), início, fim, status e observações\n\n"
           "Exclusão em cascata e índices mantêm o banco íntegro e rápido.",
           size=12.5, color=INK, leading=1.22, space_after=4)
    card(c, 7.25, yy, 5.53, 2.75, fill=NAVY, line=None)
    c.text(7.53, yy + 0.20, 5.0, 0.42, "A regra que evita conflito (SQL no Room)", size=14.5, bold=True, color=WHITE)
    c.text(7.53, yy + 0.70, 5.0, 1.0, "novoInício < fimExistente\nAND novoFim > inícioExistente",
           size=14, bold=True, color=SKY, leading=1.2, mono=True)
    c.text(7.53, yy + 1.65, 5.0, 0.95, "Ignora agendamentos CANCELADOS e o próprio registro quando ele está sendo editado.",
           size=12, color=SKY2, leading=1.2)

def s08(c):
    c.begin(BG)
    c.note("TEMPO: 7:45–9:00 | APRESENTADOR(A) 3 (uma etapa por integrante, se preferirem)\n"
           "• Conte a evolução em 4 frentes: primeiro telas/navegação, depois cadastros com banco, depois agendamento com conflito, por fim lembretes/atalhos/documentação.\n"
           "• HONESTIDADE SOBRE O GIT: o README descreve essas 4 etapas como commits planejados, mas o histórico público atual no GitHub está consolidado em um único envio ('config_system'). Falem 'quatro etapas/frentes de desenvolvimento', não 'quatro commits publicados'.\n"
           "• Essa consolidação é um dos aprendizados que retomamos no slide 13.")
    y = header(c, 8, "PROCESSO", "A evolução em 4 etapas", "Cada etapa deixou o app utilizável antes de partir para a próxima")
    c.line(0.95, y + 0.42, 12.45, y + 0.42, SKY, 3)
    etapas = [
        ("1", "Interface e navegação", "Design system em Jetpack Compose (Material 3) e navegação entre Agenda, Clientes e Serviços.", PRIMARY),
        ("2", "Clientes e serviços", "Cadastros completos (CRUD) com persistência no Room Database.", PRIMARY),
        ("3", "Agendamentos", "Calendário/data e hora, cálculo de fim pela duração e verificação de conflito.", PURPLE),
        ("4", "Lembretes e acabamento", "Notificações locais, atalhos do ícone do app e documentação (README).", PURPLE),
    ]
    for i, (n, t, b, acc) in enumerate(etapas):
        x = 0.55 + i * 3.13
        num_badge(c, x + 1.20, y + 0.14, 0.56, n, fill=acc, size=17)
        card(c, x, y + 0.92, 2.96, 2.55)
        c.text(x + 0.22, y + 1.14, 2.52, 0.62, t, size=15, bold=True, color=NAVY, align="c", leading=1.1)
        c.text(x + 0.22, y + 1.80, 2.52, 1.55, b, size=12, color=MUTED, align="c", leading=1.2)
    c.rect(0.55, y + 3.75, 12.23, 0.85, ORANGE_BG, radius=0.14)
    c.text(0.85, y + 3.75, 11.65, 0.85,
           "Registro de transparência: no GitHub, essas etapas estão hoje consolidadas em um envio principal — melhorar a granularidade dos commits entrou para a nossa lista de correções.",
           size=12.5, bold=True, color=(150, 66, 0), align="c", anchor="m")

def s09(c):
    c.begin(BG)
    c.note("TEMPO: 9:00–10:45 | APRESENTADOR(A) 4 (aprox. 35s por desafio)\n"
           "1) CONFLITO: o erro clássico seria comparar só o horário de início. A duração muda por serviço, então comparamos intervalos [início, fim) no SQL.\n"
           "2) LEMBRETE: notificação agendada esbarra nas regras novas do Android — permissão em tempo de execução (13+), alarmes exatos (12+) e alarmes perdidos ao reiniciar. Temos fallback e um receiver de boot.\n"
           "3) PORTABILIDADE: ao revisar o repositório, descobrimos que ele só compilava na máquina de origem (local.properties com caminho C:\\Users\\CLIENTE versionado, .gitignore genérico, SDK/JDK muito recentes). Diagnóstico feito; correção planejada no slide 12.")
    y = header(c, 9, "DESAFIOS", "Os 3 principais desafios", "O que travou, como enfrentamos e o que ficou diagnosticado")
    desafios = [
        ("1", "Impedir sobreposição de horários", RED,
         "Serviços têm durações diferentes: comparar só o início deixa conflitos passarem.",
         "Guardamos início e fim de cada atendimento e testamos a sobreposição de intervalos no SQL, ignorando cancelados."),
        ("2", "Lembrar mesmo com o app fechado", ORANGE,
         "Android restringe alarmes exatos (12+) e exige permissão de notificação (13+); reiniciar o celular apaga alarmes.",
         "AlarmManager 15 min antes com fallback quando não há permissão de exato + BootReceiver para reagendar após reiniciar."),
        ("3", "Fazer rodar fora da máquina de origem", PURPLE,
         "Diagnóstico: local.properties apontando para C:\\Users\\CLIENTE\\... foi versionado; .gitignore não era de Android; SDK 37/JDK 25 são exigências recentes.",
         "Correção planejada: .gitignore de Android, remover arquivos locais/build do Git e testar a partir de um clone limpo (slide 12)."),
    ]
    for i, (n, t, acc, prob, resp) in enumerate(desafios):
        x = 0.55 + i * 4.18
        card(c, x, y, 3.92, 4.55)
        c.rect(x, y, 3.92, 0.10, acc, radius=0.05)
        num_badge(c, x + 0.26, y + 0.30, 0.52, n, fill=acc)
        c.text(x + 0.92, y + 0.28, 2.80, 0.85, t, size=15.5, bold=True, color=NAVY, leading=1.12, anchor="m")
        c.text(x + 0.26, y + 1.22, 3.40, 0.3, "O problema", size=11.5, bold=True, color=acc)
        c.text(x + 0.26, y + 1.52, 3.40, 1.35, prob, size=12, color=INK, leading=1.2)
        c.line(x + 0.26, y + 2.78, x + 3.66, y + 2.78, BORDER, 1)
        c.text(x + 0.26, y + 2.90, 3.40, 0.3, "Nossa resposta", size=11.5, bold=True, color=PRIMARY)
        c.text(x + 0.26, y + 3.20, 3.40, 1.25, resp, size=12, color=MUTED, leading=1.2)

def s10(c):
    c.begin(BG)
    c.note("TEMPO: 10:45–11:45 | APRESENTADOR(A) 4\n"
           "• Liste o que está comprovado no código/demonstração, sem exageros: CRUD offline, agenda reativa, bloqueio de conflito, busca/filtros, lembrete e atalhos.\n"
           "• Use os números da direita para dar escala: 4 telas, 3 entidades, 4 status, lembrete de 15 min, Android 7+ (API 24).\n"
           "• Reforce: funcionalidades do tema estão cobertas — o gap é de formato/plataforma, próximo slide.")
    y = header(c, 10, "RESULTADOS", "O que funcionou bem", "Resultados que conseguimos Demonstrar")
    itens = [
        "Cadastros de clientes e serviços completos, salvos offline no aparelho",
        "Agenda reativa: salvou/alterou, a lista atualiza na hora (Flow/StateFlow)",
        "Conflito de horário bloqueado antes de salvar — sem dupla reserva",
        "Busca em tempo real e filtros por status do atendimento",
        "Lembrete automático 15 min antes, inclusive após reiniciar o celular",
        "Atalhos ao pressionar o ícone do app (App Shortcuts)",
    ]
    yy = y
    for it in itens:
        card(c, 0.55, yy, 7.55, 0.70)
        c.oval(0.78, yy + 0.17, 0.38, 0.38, GREEN)
        c.text(0.78, yy + 0.18, 0.38, 0.36, "✓", size=13, bold=True, color=WHITE, align="c", anchor="m")
        c.text(1.32, yy, 6.6, 0.70, it, size=13, color=INK, leading=1.15, anchor="m")
        yy += 0.82
    stats = [("4", "telas principais"), ("3", "entidades no banco"), ("4", "status de atendimento"),
             ("15 min", "antecedência do lembrete"), ("API 24+", "Android 7.0 ou superior"), ("100%", "funciona offline")]
    for i, (v, rot) in enumerate(stats):
        x = 8.40 + (i % 2) * 2.22
        yyy = y + (i // 2) * 1.65
        card(c, x, yyy, 2.08, 1.50, fill=NAVY if i % 2 == 0 else PRIMARY, line=None)
        c.text(x + 0.15, yyy + 0.16, 1.78, 0.55, v, size=23, bold=True, color=WHITE, align="c")
        c.text(x + 0.15, yyy + 0.80, 1.78, 0.55, rot, size=11, color=SKY2, align="c", leading=1.12)

def s11(c):
    c.begin(BG)
    c.note("TEMPO: 11:45–12:45 | APRESENTADOR(A) 1\n"
           "• Este é o slide mais importante da defesa: comparem requisito a requisito, sem tentar 'vender' equivalência como cumprimento.\n"
           "• ATENDE: tema/funções do Grupo 2 e código no GitHub. NÃO ATENDE: manifest.json, Service Worker, hospedagem/link público, armazenamento web e Lighthouse (sem URL, não há como medir).\n"
           "• Frase de fechamento sugerida: 'equivalência funcional não substitui os artefatos que a especificação cobra — por isso nosso próximo passo é portar para web.'\n"
           "• Se o professor tiver autorizado o Android nativo previamente, adaptem este slide e mencionem a autorização aqui.")
    y = header(c, 11, "DIAGNÓSTICO PWA", "Leitura honesta: onde estamos frente ao PWA",
               "Comparativo direto com a especificação da atividade")
    rows = [
        ("Tema do Grupo 2 (clientes, horários, status, agenda)", "Funções implementadas no app", "ATENDE", GREEN_BG, GREEN),
        ("Código versionado no GitHub", "Repositório publicado", "ATENDE", GREEN_BG, GREEN),
        ("manifest.json + Service Worker (offline web)", "AndroidManifest.xml + banco Room (offline nativo)", "NÃO ATENDE", RED_BG, RED),
        ("Hospedagem (Vercel/Netlify/Pages) + link público", "Sem deploy — roda via Android Studio/APK", "NÃO ATENDE", RED_BG, RED),
        ("LocalStorage/IndexedDB ou Firebase/Supabase", "Room (SQLite) — equivalente funcional, outra tecnologia", "NÃO ATENDE", RED_BG, RED),
        ("Responsivo web + notas do Lighthouse", "App Android, sem URL para auditar", "NÃO ATENDE", RED_BG, RED),
    ]
    c.text(0.75, y, 5.5, 0.35, "Requisito da atividade", size=12, bold=True, color=MUTED)
    c.text(6.35, y, 4.4, 0.35, "Situação no projeto", size=12, bold=True, color=MUTED)
    c.text(10.95, y, 1.7, 0.35, "Veredito", size=12, bold=True, color=MUTED, align="c")
    yy = y + 0.42
    for req, sit, ver, b, t in rows:
        card(c, 0.55, yy, 12.23, 0.58, fill=WHITE)
        c.text(0.75, yy, 5.45, 0.58, req, size=12.5, bold=True, color=NAVY, anchor="m", leading=1.1)
        c.text(6.35, yy, 4.45, 0.58, sit, size=12, color=INK, anchor="m", leading=1.1)
        pill(c, 10.95, yy + 0.12, 1.65, 0.35, ver, b, t, size=10.5)
        yy += 0.68
    c.rect(0.55, yy + 0.08, 12.23, 0.72, NAVY, radius=0.14)
    c.text(0.85, yy + 0.08, 11.65, 0.72, "Conclusão: a solução é funcional, mas equivalência nativa não gera os artefatos cobrados — manifest, Service Worker, URL pública e Lighthouse.",
           size=13, bold=True, color=WHITE, align="c", anchor="m")

def s12(c):
    c.begin(BG)
    c.note("TEMPO: 12:45–13:45 | APRESENTADOR(A) 2\n"
           "• Apresente como plano, não como feito. Coluna A corrige o repo atual (rápido). Coluna B é o trabalho de portar para PWA reaproveitando telas, regras e modelo de dados — o código Kotlin não roda no navegador.\n"
           "• No PWA: IndexedDB guarda clientes/serviços/agendamentos; a regra de conflito é a mesma lógica de intervalos; SW faz cache offline; deploy gratuito gera o link público; Lighthouse mede Performance/Acessibilidade/Boas práticas/SEO.\n"
           "• A ordem importa: primeiro PWA instalável mínimo, depois refinamento e metas do Lighthouse.")
    y = header(c, 12, "PRÓXIMOS PASSOS", "O que faríamos a seguir", "De um app Android funcional a um PWA publicável e mensurável")
    card(c, 0.55, y, 5.95, 3.55)
    pill(c, 0.85, y + 0.24, 0.55, 0.42, "A", NAVY, WHITE, size=14)
    c.text(1.55, y + 0.24, 4.7, 0.45, "Corrigir o repositório atual", size=16.5, bold=True, color=NAVY, anchor="m")
    c.text(0.85, y + 0.85, 5.35, 2.55,
           "• .gitignore próprio de Android\n"
           "• Retirar do Git: local.properties, build/, .gradle/ e .idea/\n"
           "• Permissão de execução no gradlew e versões coerentes no README\n"
           "• Commits pequenos por etapa e teste a partir de clone limpo",
           size=13.5, color=INK, leading=1.22, space_after=5)
    card(c, 6.85, y, 5.93, 3.55, fill=SKY_SOFT, line=SKY)
    pill(c, 7.15, y + 0.24, 0.55, 0.42, "B", PRIMARY, WHITE, size=14)
    c.text(7.85, y + 0.24, 4.7, 0.45, "Portar a solução para PWA", size=16.5, bold=True, color=NAVY, anchor="m")
    c.text(7.15, y + 0.85, 5.35, 2.55,
           "• Recriar as telas em web responsiva (mobile-first)\n"
           "• IndexedDB com o mesmo modelo e a mesma regra de conflito\n"
           "• manifest.json (ícones e atalhos) + Service Worker (cache offline e notificações)\n"
           "• Deploy gratuito (GitHub Pages/Vercel) e metas do Lighthouse",
           size=13.5, color=INK, leading=1.22, space_after=5)
    yy = y + 3.85
    steps = ["1  Higiene do repo", "2  MVP web", "3  Camada PWA", "4  Deploy + Lighthouse"]
    for i, s in enumerate(steps):
        x = 0.55 + i * 3.13
        pill(c, x, yy, 2.88, 0.55, s, NAVY if i % 2 == 0 else PRIMARY, WHITE, size=12.5)
        if i < 3:
            c.text(x + 2.88, yy, 0.26, 0.55, "→", size=14, bold=True, color=MUTED, align="c", anchor="m")

def s13(c):
    c.begin(BG)
    c.note("TEMPO: 13:45–14:30 | APRESENTADOR(A) 3\n"
           "• Personalizem: citem um exemplo concreto do grupo para cada lição (quem viveu o quê).\n"
           "• A lição 1 é a maior: mergulhamos na tecnologia que dominávamos antes de reler a especificação. A lição 4 fecha o ciclo: avaliar como o professor avalia (clone limpo, link) teria revelado os problemas cedo.\n"
           "• Experiência da equipe: trabalhar por frentes (interface, dados, agendamento, sistema/documentação) ensinou integração — cada entrega dependia da anterior.\n"
           "TRANSIÇÃO: 'Para fechar...'")
    y = header(c, 13, "LIÇÕES", "O que levamos desta atividade", "Aprendizados técnicos e de processo da equipe")
    licoes = [
        ("1", "Requisito antes da tecnologia", "A especificação define o formato (PWA). Escolher a ferramenta antes de relê-la gerou o nosso maior desvio."),
        ("2", "Offline é arquitetura", "Persistência local, estado reativo e regras no banco foram decisões de projeto — não um detalhe no fim."),
        ("3", "Versionamento também é engenharia", ".gitignore correto, nada de caminho local no Git e commits pequenos valem tanto quanto código funcionando."),
        ("4", "Testar como o avaliador testa", "Clonar em máquina limpa e abrir pelo link público revelaria cedo os problemas de ambiente e de entrega."),
    ]
    for i, (n, t, b) in enumerate(licoes):
        x = 0.55 + (i % 2) * 6.20
        yy = y + (i // 2) * 1.72
        card(c, x, yy, 5.95, 1.56)
        num_badge(c, x + 0.26, yy + 0.26, 0.5, n, fill=PURPLE if i % 2 else PRIMARY)
        c.text(x + 0.92, yy + 0.20, 4.8, 0.42, t, size=15.5, bold=True, color=NAVY)
        c.text(x + 0.92, yy + 0.64, 4.80, 0.82, b, size=12, color=MUTED, leading=1.18)
    yy = y + 3.62
    c.rect(0.55, yy, 12.23, 1.02, NAVY, radius=0.14)
    c.text(0.85, yy, 11.65, 1.02,
           "Experiência da equipe: dividir o trabalho em frentes (interface, dados, agendamento e sistema/documentação) mostrou que integração e leitura crítica do requisito importam tanto quanto programar.",
           size=13.5, bold=True, color=WHITE, align="c", anchor="m")

def s14(c):
    c.begin(NAVY)
    c.note("TEMPO: 14:30–15:00 | TODOS (fecha o(a) Apresentador(a) 4)\n"
           "• Recapitule em 2 frases: entregamos uma agenda de serviços funcional e offline; reconhecemos que ela ainda não é o PWA hospedado da especificação e apresentamos o caminho.\n"
           "• Aponte o repositório e convide para perguntas. Deixem o app aberto no emulador/celular caso peçam para ver alguma função.\n"
           "• Perguntas prováveis: por que não web desde o início? como funciona o conflito? como a notificação dispara com o app fechado? — respostas nos slides 7, 9 e 11.")
    c.oval(10.2, -2.4, 6.0, 6.0, NAVY2)
    c.oval(-2.2, 4.9, 4.8, 4.8, PRIMARY_D)
    pill(c, 0.65, 0.60, 2.55, 0.42, "ENCERRAMENTO", None, SKY, size=10.5, line=PRIMARY)
    c.text(0.65, 1.25, 8.5, 1.1, "Obrigado!", size=52, bold=True, color=WHITE)
    c.text(0.65, 2.35, 9.6, 0.6, "Perguntas?", size=26, bold=True, color=SKY)
    c.text(0.65, 3.15, 10.6, 1.0,
           "Entregamos uma base funcional e offline de agenda de serviços.\nO próximo passo é transformá-la, de fato, em um PWA publicável, instalável e mensurável.",
           size=16, color=WHITE, leading=1.25)
    card(c, 0.65, 4.45, 11.95, 1.30, fill=WHITE, line=None)
    c.text(0.95, 4.62, 3.2, 0.4, "Repositório no GitHub", size=12.5, bold=True, color=MUTED)
    c.text(0.95, 5.02, 6.9, 0.5, "github.com/LuizFelip908/Agenda-de-Servicos", size=16, bold=True, color=NAVY)
    pill(c, 8.15, 4.62, 2.05, 0.42, "✓ Código publicado", GREEN_BG, GREEN, size=10.5)
    pill(c, 10.30, 4.62, 2.05, 0.42, "Link público: a portar", ORANGE_BG, ORANGE, size=10.5)
    c.text(8.15, 5.14, 4.20, 0.45, "Demonstração: app Android (emulador/celular)", size=11, color=MUTED)
    c.text(0.65, 6.15, 12.0, 0.8,
           "Grupo 2 — [Nome 1]  •  [Nome 2]  •  [Nome 3]  •  [Nome 4]     |     Agenda de Serviços — barbearia & oficina",
           size=13, color=SKY2)

SLIDES = [s01, s02, s03, s04, s05, s06, s07, s08, s09, s10, s11, s12, s13, s14]

def build(canvas):
    for fn in SLIDES:
        fn(canvas)
        canvas.end()
    canvas.save()

if __name__ == "__main__":
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Agenda-de-Servicos-Grupo2")
    build(PptxCanvas(base + ".pptx"))
    build(PdfCanvas(base + ".pdf"))
    print("Gerados:", base + ".pptx", "e", base + ".pdf")
